from __future__ import annotations

import html
import os
import re
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path

import pymupdf
from docx import Document
from pptx import Presentation


TEXT_EXTENSIONS = {
    ".txt", ".md", ".rst", ".csv", ".json", ".jsonl", ".yaml", ".yml",
    ".toml", ".xml", ".log", ".py", ".js", ".ts", ".java", ".go", ".rs",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | IMAGE_EXTENSIONS | {
    ".pdf", ".docx", ".pptx", ".html", ".htm"
}


@dataclass(frozen=True)
class Section:
    text: str
    locator: str


def discover(path: Path) -> list[Path]:
    files = [path] if path.is_file() else [item for item in path.rglob("*") if item.is_file()]
    return sorted(file for file in files if file.suffix.lower() in SUPPORTED_EXTENSIONS)


def _ocr(path: Path) -> str:
    try:
        result = subprocess.run(
            [
                "tesseract",
                str(path),
                "stdout",
                "-l",
                os.getenv("LOCAL_RAG_OCR_LANGUAGE", "eng"),
            ],
            check=True,
            capture_output=True,
            text=True,
        )
    except FileNotFoundError as exc:
        raise RuntimeError("OCR requires Tesseract: https://tesseract-ocr.github.io/") from exc
    except subprocess.CalledProcessError as exc:
        raise RuntimeError(exc.stderr.strip() or f"OCR failed for {path}") from exc
    return result.stdout


def extract(path: Path) -> list[Section]:
    suffix = path.suffix.lower()
    if suffix in TEXT_EXTENSIONS:
        return [Section(path.read_text(encoding="utf-8", errors="replace"), "file")]
    if suffix in {".html", ".htm"}:
        raw = path.read_text(encoding="utf-8", errors="replace")
        text = re.sub(r"<script.*?</script>|<style.*?</style>", " ", raw, flags=re.I | re.S)
        return [Section(html.unescape(re.sub(r"<[^>]+>", " ", text)), "file")]
    if suffix in IMAGE_EXTENSIONS:
        return [Section(_ocr(path), "image")]
    if suffix == ".docx":
        text = "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
        return [Section(text, "document")]
    if suffix == ".pptx":
        presentation = Presentation(path)
        return [
            Section(
                "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text")),
                f"slide {number}",
            )
            for number, slide in enumerate(presentation.slides, 1)
        ]
    if suffix == ".pdf":
        sections = []
        with pymupdf.open(path) as document:
            for number, page in enumerate(document, 1):
                text = page.get_text().strip()
                if not text:
                    with tempfile.NamedTemporaryFile(suffix=".png") as image:
                        page.get_pixmap(matrix=pymupdf.Matrix(2, 2)).save(image.name)
                        text = _ocr(Path(image.name)).strip()
                sections.append(Section(text, f"page {number}"))
        return sections
    raise ValueError(f"unsupported file: {path}")


def chunks(text: str, size: int = 1200, overlap: int = 200) -> list[str]:
    if size <= overlap or overlap < 0:
        raise ValueError("chunk size must be greater than overlap")
    clean = re.sub(r"\s+", " ", text).strip()
    if not clean:
        return []
    result = []
    start = 0
    while start < len(clean):
        end = min(start + size, len(clean))
        if end < len(clean):
            boundary = clean.rfind(" ", start + size // 2, end)
            end = boundary if boundary > start else end
        result.append(clean[start:end].strip())
        if end == len(clean):
            break
        start = end - overlap
    return result
